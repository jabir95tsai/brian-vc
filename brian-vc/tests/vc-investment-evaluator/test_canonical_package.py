#!/usr/bin/env python3
"""End-to-end state/gate test through the F1 canonical package."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import unittest
import uuid
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "skills" / "vc-investment-evaluator" / "scripts"
RUNNER = SCRIPTS / "evaluator_runner.py"
PREPARE = SCRIPTS / "prepare_workbook_input.py"
ASSEMBLE = SCRIPTS / "assemble_canonical_package.py"
FIXTURE = Path(__file__).with_name("fixture_evaluator_case.json")
MODULES_TO_E_GATE = ["A1", "A2", "B1", "B3", "B4", "C1", "C2", "C3", "C4", "D1", "D2", "D3", "E1", "E2", "E3"]


class CanonicalPackageTests(unittest.TestCase):
    def setUp(self) -> None:
        self.case_dir = Path(__file__).with_name(".canonical-test-runtime") / uuid.uuid4().hex
        self.case_dir.mkdir(parents=True)

    def tearDown(self) -> None:
        shutil.rmtree(self.case_dir, ignore_errors=True)

    def run_cmd(self, *args: str) -> subprocess.CompletedProcess[str]:
        return subprocess.run([sys.executable, *args], text=True, encoding="utf-8", capture_output=True, check=False)

    def exercise_freeze(self, mode="full", tamper=None, deliver=False) -> None:
        init = self.run_cmd(str(RUNNER), "init", str(self.case_dir), "--case-id", "20260815_E2E", "--mode", mode)
        self.assertEqual(init.returncode, 0, init.stderr)
        evidence_dir = self.case_dir / "evidence"
        evidence_dir.mkdir()
        for module in MODULES_TO_E_GATE:
            artifact = evidence_dir / f"{module}.md"
            artifact.write_text(f"# {module}\nverified fixture evidence\n", encoding="utf-8")
            # D1 and E2 owe a named fact at completion (pipeline contract s.2).
            key_facts = {
                "D1": "peer_list_source=auto",
                "E2": "redteam_handoff=RedTeam 提出 5 個反對理由，主要風險點為 A、B、C，"
                      "GP 決策框架已留白供填入。",
            }
            evidence_args = ["--evidence", str(artifact)]
            if module in key_facts:
                evidence_args += ["--evidence", key_facts[module]]
            status = "complete"
            if mode == "blocked" and module in ("D2", "D3"):
                status = "blocked"
                evidence_args += ["--reason", "synthetic missing Term Sheet",
                                  "--evidence", "blocked_as_designed=missing_transaction_terms"]
            done = self.run_cmd(str(RUNNER), "set", str(self.case_dir), module, status, *evidence_args, "--artifact", str(artifact))
            self.assertEqual(done.returncode, 0, f"{module}: {done.stderr or done.stdout}")
            if module == "B1":
                b2_evidence = evidence_dir / "B2-not-applicable.md"
                b2_evidence.write_text("DocumentIndex contains no prospectus", encoding="utf-8")
                b2 = self.run_cmd(str(RUNNER), "set", str(self.case_dir), "B2", "not_applicable", "--evidence", str(b2_evidence), "--reason", "no prospectus in DocumentIndex")
                self.assertEqual(b2.returncode, 0, b2.stderr or b2.stdout)
        prepared = self.case_dir / "prepared_case.json"
        fixture = json.loads(FIXTURE.read_text(encoding="utf-8"))
        fixture["mode"] = mode
        if mode == "blocked":
            for key in ("investment", "pre_money", "post_money", "price_per_share"):
                fixture["deal"][key] = None
            fixture["deal"]["blocked_reason"] = "synthetic missing Term Sheet"
            fixture.pop("deck", None)
        source = self.case_dir / "synthetic-input.json"
        source.write_text(json.dumps(fixture, ensure_ascii=False), encoding="utf-8")
        prep = self.run_cmd(str(PREPARE), str(source), str(prepared))
        self.assertEqual(prep.returncode, 0, prep.stderr)
        # E2 supplies the handoff before F1; a generated scaffold alone is not
        # an expert review. Use the existing synthetic fixture's E2 text.
        payload = json.loads(prepared.read_text(encoding="utf-8"))
        payload["deck"]["redteam_handoff"] = json.loads(FIXTURE.read_text(encoding="utf-8"))["deck"]["redteam_handoff"]
        prepared.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        if tamper == "hash":
            (evidence_dir / "D2.md").write_text("changed synthetic evidence", encoding="utf-8")
        if tamper in ("mode", "matrix"):
            payload = json.loads(prepared.read_text(encoding="utf-8"))
            if tamper == "mode":
                payload["mode"] = "full"
            else:
                payload["return_matrix"]["irr_rows"] = [[0]]
            prepared.write_text(json.dumps(payload), encoding="utf-8")
        assembled = self.run_cmd(str(ASSEMBLE), str(self.case_dir), str(prepared))
        if tamper:
            self.assertNotEqual(assembled.returncode, 0)
            expected_error = {"hash": "changed", "mode": "mode differs", "matrix": "calculation refusals"}[tamper]
            self.assertIn(expected_error, assembled.stderr)
            self.assertFalse((self.case_dir / "outputs" / "20260815_E2E_ContextPackage.json").exists())
            return
        self.assertEqual(assembled.returncode, 0, assembled.stderr or assembled.stdout)
        report = json.loads(assembled.stdout)
        package = Path(report["context_package"])
        self.assertTrue(package.is_file())
        state = json.loads((self.case_dir / ".vc-evaluator" / "artifact-manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(state["modules"]["F1"]["status"], "complete")
        status = self.run_cmd(str(RUNNER), "status", str(self.case_dir), "--json")
        self.assertEqual(json.loads(status.stdout)["gates"]["E_GATE"], "complete")
        if mode == "blocked":
            frozen = json.loads(package.read_text(encoding="utf-8"))
            self.assertEqual(frozen["gates"]["D_GATE"], "blocked")
            self.assertNotIn("irr_rows", frozen["content"]["return_matrix"])
            self.assertIsNone(frozen["content"]["deal"]["investment"])
            for key in ("D2", "D3"):
                self.assertEqual(state["modules"][key]["status"], "blocked")
                refused = self.run_cmd(str(RUNNER), "set", str(self.case_dir), key, "complete",
                                       "--evidence", "bogus", "--artifact", str(package))
                self.assertNotEqual(refused.returncode, 0)
        verified = self.run_cmd(str(RUNNER), "verify", str(self.case_dir), "--json")
        self.assertEqual(verified.returncode, 0, verified.stdout)
        if deliver:
            # Exercise the actual delivery CLI and runner through F_GATE. The
            # workbook audit and visual report are synthetic boundary inputs,
            # NOT evidence of real rendering, workbook computation or visual QA.
            from pptx import Presentation
            from pptx.util import Inches
            decks = []
            for variant in ("exec", "full"):
                prs = Presentation()
                for index in range(18):
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    box = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(9), Inches(1))
                    box.text_frame.text = "SYNTHETIC gate test: 同業比較 IRR blocked 財測 blocked 團隊待補 歷史財務 技術 產業 風險矩陣 RedTeam 補件"
                    if index == 0:
                        for offset in (1, 4):
                            table = slide.shapes.add_table(9, 2, Inches(0), Inches(offset), Inches(8), Inches(2)).table
                            for row in table.rows:
                                for cell in row.cells:
                                    cell.text = "synthetic evidence"
                path = self.case_dir / f"{variant}.pptx"
                prs.save(path)
                decks.append(path)
            workbooks = [self.case_dir / "factbase.xlsx", self.case_dir / "model.xlsx"]
            for path in workbooks:
                path.write_bytes(b"SYNTHETIC audit-boundary placeholder; not a workbook")
            audit = self.case_dir / "audit.json"
            audit.write_text(json.dumps({"factbase": {"formula_error_count": 0}, "model": {
                "formula_error_count": 0, "formula_count_in_base_forecast": 0,
                "model_checks": "BLOCKED_AS_DESIGNED"}}), encoding="utf-8")
            visual = self.case_dir / "synthetic-visual-report.json"
            visual.write_text(json.dumps({"status": "pass", "reviewed_files": [str(x.resolve()) for x in decks],
                "note": "Synthetic validator input; no actual visual inspection"}), encoding="utf-8")
            delivered = self.run_cmd(str(SCRIPTS / "verify_and_record_delivery.py"), str(self.case_dir),
                "--context-package", str(package), "--factbase", str(workbooks[0]),
                "--financial-model", str(workbooks[1]), "--executive", str(decks[0]),
                "--full-critical", str(decks[1]), "--workbook-audit", str(audit),
                "--visual-qa-report", str(visual), "--mode", mode)
            self.assertEqual(delivered.returncode, 0, delivered.stderr)
            self.assertEqual(json.loads(delivered.stdout)["F_GATE"], "complete")
            report = json.loads(self.run_cmd(str(RUNNER), "status", str(self.case_dir), "--json").stdout)
            self.assertEqual(report["gates"]["D_GATE"], "blocked")
            self.assertEqual(report["modules"]["D3"]["status"], "blocked")

    def test_f1_freeze_records_context_package(self) -> None:
        self.exercise_freeze()

    def test_degraded_positive_path(self) -> None:
        self.exercise_freeze("degraded")

    def test_missing_terms_freezes_facts_without_completing_transaction_modules(self) -> None:
        self.exercise_freeze("blocked")

    def test_blocked_delivery_cli_reaches_f_gate_with_synthetic_audit_inputs(self) -> None:
        self.exercise_freeze("blocked", deliver=True)

    def test_blocked_freeze_rejects_stale_evidence(self) -> None:
        self.exercise_freeze("blocked", "hash")

    def test_blocked_freeze_rejects_mode_mismatch(self) -> None:
        self.exercise_freeze("blocked", "mode")

    def test_blocked_freeze_rejects_injected_returns(self) -> None:
        self.exercise_freeze("blocked", "matrix")


if __name__ == "__main__":
    unittest.main(verbosity=2)
